"""“이 PPT의 모든 숫자와 출처를 검증해줘” — document fact-check mode."""
import io

from pptx import Presentation
from pptx.util import Inches

from app.jobs.queue import drain
from app.pipeline.doc_verify import judge


def _deck() -> bytes:
    prs = Presentation()
    for title, body in [("시장", "국내 AI 헬스케어 시장 규모는 2024년 1조 2,000억원으로 집계되었다."),
                        ("경쟁", "국내 AI 헬스케어 시장 규모는 2024년 3조원으로 집계되었다."),
                        ("성과", "우리 서비스의 누적 사용자는 2025년 50만명을 돌파했다.")]:
        s = prs.slides.add_slide(prs.slide_layouts[5])
        s.shapes.title.text = title
        s.shapes.add_textbox(Inches(1), Inches(2), Inches(8), Inches(1)).text_frame.text = body
    b = io.BytesIO()
    prs.save(b)
    return b.getvalue()


def test_judge_rules():
    assert judge("시장 규모는 2024년 1조 2,000억원이다.", "국내 시장 규모는 2024년 1조 2,000억원으로 집계되었다.")[0] == "verified"
    assert judge("국내 시장 규모는 2024년 3조원이다.", "국내 시장 규모는 2024년 1조 2,000억원으로 집계되었다.")[0] == "contradicted"
    assert judge("사용자는 50만명이다.", "날씨가 좋다.")[0] == "unverified"


def test_deck_fact_check_end_to_end(client):
    p = client.post("/api/projects", json={"name": "IR 덱 검증"}).json()
    up = client.post("/api/documents", files={"file": ("ir_deck.pptx", _deck())}, data={"project_id": str(p["id"])}).json()
    t = client.post("/api/tasks", json={"request": "이 PPT의 모든 숫자와 출처를 검증해줘", "project_id": p["id"]}).json()
    drain()
    t = client.get(f"/api/tasks/{t['id']}").json()
    assert t["status"] == "WAITING_USER", t["error"]
    assert t["plan"]["mode"] == "verify_document" and t["plan"]["target_document_id"] == up["id"]
    claims = {c["text"]: c for c in client.get(f"/api/claims?task_id={t['id']}").json()}
    ok = claims["국내 AI 헬스케어 시장 규모는 2024년 1조 2,000억원으로 집계되었다."]
    bad = claims["국내 AI 헬스케어 시장 규모는 2024년 3조원으로 집계되었다."]
    none = claims["우리 서비스의 누적 사용자는 2025년 50만명을 돌파했다."]
    assert ok["verification_status"] == "verified" and ok["source_id"] and "1조 2,000억원" in ok["supporting_quote"]
    assert bad["verification_status"] == "contradicted"
    assert none["verification_status"] == "unverified" and none["source_id"] is None
    rs = t["result_summary"]
    assert rs["verdicts"]["verified"] >= 1 and rs["qc"]["passed"]
    import docx

    text = "\n".join(p.text for p in docx.Document(rs["files"]["docx"]).paragraphs)
    assert "수정이 필요한 항목" in text and "3조원" in text


def test_explicit_document_id(client):
    up = client.post("/api/documents", files={"file": ("memo.md", "# 메모\n\n국내 AI 헬스케어 시장 규모는 2024년 1조 2,000억원이다.\n".encode())}).json()
    t = client.post("/api/tasks", json={"request": "이 메모 확인해줘", "document_id": up["id"]}).json()
    assert client.post("/api/tasks", json={"request": "x 검증", "document_id": 999999}).status_code == 400
    drain()
    t = client.get(f"/api/tasks/{t['id']}").json()
    assert t["plan"]["mode"] == "verify_document" and t["result_summary"]["claims_checked"] >= 1
