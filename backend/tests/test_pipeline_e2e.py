"""End-to-end: web research → sources → claims → verification → critic → report/PPT/XLSX → QC → 결재 → final.
Runs through the public API with the research snapshot (search index + mirrored originals)."""
import zipfile
from pathlib import Path

import pytest

from app.jobs.queue import drain

REQUEST = "AI 헬스케어 시장을 조사해서 새로운 사업 아이디어를 만들고 투자자용 PPT까지 만들어줘"


@pytest.fixture(scope="module")
def finished_task(client):
    r = client.post("/api/tasks", json={"request": REQUEST, "priority": 1})
    assert r.status_code == 201
    tid = r.json()["id"]
    assert drain() >= 1
    return client.get(f"/api/tasks/{tid}").json()


def test_task_reaches_ceo_approval(finished_task):
    t = finished_task
    assert t["status"] == "WAITING_USER", t.get("error")
    assert t["progress"] == 100
    names = [s["name"] for s in t["steps"]]
    for step in ("요구사항 분석", "킥오프 회의", "웹 조사·원문 확보", "근거 추출", "사실관계 검증", "비판 검토·리뷰 회의",
                 "분석·전략·재무", "보고서 작성", "발표자료 제작", "스프레드시트", "품질 검사·결재"):
        assert step in names
    assert all(s["status"] == "done" for s in t["steps"])


def test_artifacts_are_real_files(finished_task):
    arts = {a["kind"]: a for a in finished_task["artifacts"]}
    assert {"report", "presentation", "spreadsheet"} <= set(arts)
    files = finished_task["result_summary"]["files"]
    for fmt in ("docx", "pptx", "xlsx"):
        assert Path(files[fmt]).exists() and Path(files[fmt]).stat().st_size > 5000
        assert zipfile.is_zipfile(files[fmt])
    assert "_v1." in files["docx"]


def test_sources_prioritised_and_failures_reported(client, finished_task):
    summary = finished_task["result_summary"]
    assert summary["sources_accessed"] >= 5
    assert any("paywalled-news.invalid" in f for f in summary["sources_failed"])  # not silently dropped
    sources = client.get(f"/api/sources?task_id={finished_task['id']}").json()
    assert sources[0]["tier"] == 1  # official sources first
    assert all(s["accessed"] or s["access_error"] for s in sources)


def test_claims_verification_outcomes(client, finished_task):
    claims = client.get(f"/api/claims?task_id={finished_task['id']}").json()
    by_text = {c["text"]: c for c in claims}
    gov = by_text["국내 AI 헬스케어 시장 규모는 2024년 1조 2,000억원으로 집계되었다."]
    blog = by_text["국내 AI 헬스케어 시장 규모는 2024년 5조원으로 집계되었다."]
    assert gov["verification_status"] == "verified" and gov["page_number"] == 1
    assert blog["verification_status"] == "contradicted"
    assert any(c["verification_status"] == "outdated" for c in claims)  # 2021 OECD report
    # the prompt-injection text in the blog never becomes evidence
    assert not any("API key" in c["text"] or "system prompt" in c["text"].lower() for c in claims)
    assert all(c["supporting_quote"] for c in claims)


def test_citation_cards(client, finished_task):
    cards = client.get(f"/api/citations/{finished_task['id']}").json()
    assert cards[0]["number"] == 1 and cards[0]["url"] and cards[0]["passages"][0]["passage"]
    assert {"title", "publisher", "date", "url", "passages"} <= set(cards[0])


def test_report_citations_match_bibliography(finished_task):
    import re

    import docx

    path = finished_task["result_summary"]["files"]["docx"]
    xml = zipfile.ZipFile(path).read("word/document.xml").decode()
    cited = {int(n) for n in re.findall(r'w:anchor="ref_(\d+)"', xml)}
    bib = {int(n) for n in re.findall(r'w:name="ref_(\d+)"', xml)}
    assert cited and cited == bib
    text = "\n".join(p.text for p in docx.Document(path).paragraphs)
    assert "[미확인]" in text  # unknowns are shown, not filled in
    assert "5조원" in text and "수치 충돌" in text  # contradiction disclosed in limitations


def test_pptx_storyline_and_chart(finished_task):
    from pptx import Presentation

    prs = Presentation(finished_task["result_summary"]["files"]["pptx"])
    kickers = [sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.isupper()]
    for k in ("PROBLEM", "EVIDENCE", "INSIGHT", "SOLUTION", "BUSINESS MODEL", "MARKET", "STRATEGY", "FINANCIALS", "ROADMAP", "CONCLUSION"):
        assert k in kickers
    charts = [sh.chart for s in prs.slides for sh in s.shapes if sh.has_chart]
    market = charts[0]
    assert list(market.plots[0].categories) == ["2023", "2024", "2028(E)"]
    assert list(market.plots[0].series[0].values) == [9500.0, 12000.0, 35000.0]


def test_qc_passed_and_counts(finished_task):
    qc = finished_task["result_summary"]["qc"]
    assert qc["passed"], [c for c in qc["checks"] if not c["passed"]]
    assert qc["stats"]["claims_checked"] >= 10


def test_office_conversation_recorded(client, finished_task):
    msgs = client.get(f"/api/office/messages?task_id={finished_task['id']}").json()
    kinds = {m["kind"] for m in msgs}
    assert {"meeting", "approval_request", "approval_stamp", "progress"} <= kinds
    meetings = finished_task["meetings"]
    assert len(meetings) >= 2 and "| Action Item | Owner | Deadline | Status |" in meetings[0]["minutes"]


def test_audit_log_covers_workflow(client, finished_task):
    actions = {a["action"] for a in client.get(f"/api/tasks/{finished_task['id']}/audit").json()}
    assert {"task_created", "search_performed", "source_downloaded", "claims_verified", "artifact_generated", "task_completed"} <= actions


def test_ceo_final_approval_promotes_to_final(client, finished_task):
    appr_id = finished_task["result_summary"]["final_approval_id"]
    r = client.post(f"/api/approvals/{appr_id}/decide", json={"approve": True, "note": "좋습니다"})
    assert r.status_code == 200, r.text
    finals = r.json()["final_files"]
    assert finals and all("/final/" in f and Path(f).exists() for f in finals)
    t = client.get(f"/api/tasks/{finished_task['id']}").json()
    assert t["status"] == "DONE"
    assert client.post(f"/api/approvals/{appr_id}/decide", json={"approve": True}).status_code == 409


def test_revision_creates_new_version_without_overwriting(client, finished_task):
    tid = finished_task["id"]
    before = {a["id"]: a["version"] for a in client.get(f"/api/artifacts?task_id={tid}").json()}
    assert client.post(f"/api/tasks/{tid}/revise", json={"feedback": "시장 섹션을 더 강조해 주세요"}).status_code == 202
    drain()
    after = client.get(f"/api/artifacts?task_id={tid}").json()
    report = next(a for a in after if a["kind"] == "report" and a["versions"][0]["format"] == "docx")
    assert report["version"] == before[report["id"]] + 1
    labels = [v["label"] for v in report["versions"]]
    assert labels[:3] == ["v1", "final", "v2"] or {"v1", "v2", "final"} <= set(labels)
    paths = [a for a in after if a["id"] == report["id"]][0]["versions"]
    assert len({v["file_name"] for v in paths}) == len(paths)  # every version is its own file


def test_download_and_export(client, finished_task):
    arts = client.get(f"/api/artifacts?task_id={finished_task['id']}&kind=spreadsheet").json()
    r = client.get(f"/api/artifacts/{arts[0]['id']}/download")
    assert r.status_code == 200 and r.content[:2] == b"PK"
    csv = client.get(f"/api/artifacts/{arts[0]['id']}/export?format=csv")
    assert csv.status_code == 200 and b"," in csv.content
    rep = next(a for a in client.get(f"/api/artifacts?task_id={finished_task['id']}&kind=report").json() if a["versions"][0]["format"] == "docx")
    md = client.get(f"/api/artifacts/{rep['id']}/export?format=md")
    assert md.status_code == 200 and "참고문헌".encode() in md.content
