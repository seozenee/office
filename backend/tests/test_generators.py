"""Real file generation tests: test.docx, test.pptx, test.xlsx are written to disk and re-opened."""
import zipfile
from datetime import datetime, timezone
from types import SimpleNamespace

import openpyxl
import pytest
from docx import Document
from pptx import Presentation

from app.generators.content import Block, CitationRegistry, ReportContent, Section, Statement
from app.generators.docx_builder import build_docx
from app.generators.export import to_markdown, to_pdf
from app.generators.pptx_builder import SlideSpec, build_pptx
from app.generators.xlsx_builder import XlsxBuilder, evaluate_workbook, verify_model_sheet
from app.tools.calc import compute_financials, default_assumptions


@pytest.fixture()
def content():
    reg = CitationRegistry()
    src = lambda i, t: SimpleNamespace(id=i, title=t, publisher="기관", publication_date="2025-01-01",  # noqa: E731
                                       url=f"https://example.go.kr/{i}", access_date=datetime.now(timezone.utc), source_type="government")
    n1, n2 = reg.cite(src(10, "정부 보고서"), page=3), reg.cite(src(20, "학술 논문"))
    assert (n1, n2, reg.cite(src(10, "정부 보고서"))) == (1, 2, 1)  # stable numbering
    return ReportContent(
        title="테스트 보고서", subtitle="부제",
        summary=[Statement("시장 규모는 10조원이다.", "FACT", [n1], footnote="원문 p.3: 시장 규모는 10조원")],
        sections=[Section("분석", [Block("paragraph", [Statement("성장 가능성이 있다.", "ANALYSIS", [n2])]),
                                   Block("table", header=["연도", "값"], rows=[["2024", "1"]], caption="표 1"),
                                   Block("chart", chart={"kind": "line", "title": "추이", "categories": ["2024", "2025"],
                                                         "series": [{"name": "값", "values": [1, 2]}]})])],
        bibliography=reg.entries)


def test_docx_generation(content, tmp_root):
    path = build_docx(content, tmp_root / "test.docx")
    d = Document(str(path))
    texts = "\n".join(p.text for p in d.paragraphs)
    assert "테스트 보고서" in texts and "참고문헌" in texts and "[1]" in texts
    assert any(p.style.name == "Heading 1" for p in d.paragraphs)
    assert len(d.tables) >= 1 and len(d.inline_shapes) == 1  # table + chart image
    z = zipfile.ZipFile(path)
    doc_xml = z.read("word/document.xml").decode()
    assert "word/footnotes.xml" in z.namelist() and "footnoteReference" in doc_xml
    assert 'w:anchor="ref_1"' in doc_xml and 'w:name="ref_1"' in doc_xml  # clickable citation → bibliography bookmark
    assert "TOC" in doc_xml and "PAGE" in z.read("word/footer1.xml").decode()


def test_docx_every_citation_has_bibliography_entry(content, tmp_root):
    import re

    path = build_docx(content, tmp_root / "cite.docx")
    xml = zipfile.ZipFile(path).read("word/document.xml").decode()
    anchors = set(re.findall(r'w:anchor="(ref_\d+)"', xml))
    bookmarks = set(re.findall(r'w:name="(ref_\d+)"', xml))
    assert anchors and anchors <= bookmarks


def test_pptx_generation(content, tmp_root):
    slides = [SlideSpec("title", "덱", subtitle="부제", notes="노트"),
              SlideSpec("bullets", "문제", kicker="Problem", statements=content.summary, notes="발표자 노트"),
              SlideSpec("chart", "시장", chart={"kind": "bar", "categories": ["a", "b"], "series": [{"name": "s", "values": [1, 2]}]}),
              SlideSpec("table", "표", header=["a", "b"], rows=[["1", "2"]]),
              SlideSpec("timeline", "로드맵", steps=[{"label": "Q1", "text": "x"}, {"label": "Q2", "text": "y"}]),
              SlideSpec("process", "프로세스", steps=[{"label": "A", "text": "x"}]),
              SlideSpec("kpi", "KPI", kpis=[{"value": "10", "label": "k"}]),
              SlideSpec("two_column", "비교", statements=content.summary, right_statements=content.summary),
              SlideSpec("sources", "출처", bibliography=content.bibliography)]
    path = build_pptx("덱", slides, tmp_root / "test.pptx")
    prs = Presentation(str(path))
    assert len(prs.slides) == 9
    assert any(sh.has_chart for sh in prs.slides[2].shapes)
    assert any(getattr(sh, "has_table", False) and sh.has_table for sh in prs.slides[3].shapes)
    assert prs.slides[1].notes_slide.notes_text_frame.text == "발표자 노트"
    all_text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "[1]" in all_text and "https://example.go.kr/10" in all_text


def test_xlsx_generation_formulas_match_python(tmp_root):
    model = compute_financials(default_assumptions())
    x = XlsxBuilder("테스트")
    x.add_summary([("항목", 1)])
    x.add_assumptions(model.assumptions)
    x.add_financial_model(model)
    x.add_table_sheet("Claims", ["id", "text"], [[1, "a"]])
    path = x.save(tmp_root / "test.xlsx")
    wb = openpyxl.load_workbook(path)
    assert {"Summary", "Assumptions", "Model", "Claims"} <= set(wb.sheetnames)
    assert str(wb["Model"]["E4"].value).startswith("=")  # live formulas, not pasted numbers
    assert wb["Model"]._charts
    assert verify_model_sheet(path, model) == []


def test_xlsx_formula_check_detects_tampering(tmp_root):
    model = compute_financials(default_assumptions())
    x = XlsxBuilder("t")
    x.add_assumptions(model.assumptions)
    x.add_financial_model(model)
    path = x.save(tmp_root / "tamper.xlsx")
    wb = openpyxl.load_workbook(path)
    wb["Model"]["E5"] = "=B5*D5*2"
    wb.save(path)
    problems = verify_model_sheet(path, model)
    assert problems and "E5" in problems[0]
    assert evaluate_workbook(path)["Model"]["E5"] == pytest.approx(model.rows[1]["revenue"] * 2, rel=1e-6)


def test_markdown_and_pdf_export(content, tmp_root):
    md = to_markdown(content)
    assert "## 참고문헌" in md and "[1]" in md and "| 연도 | 값 |" in md
    pdf = to_pdf(content, tmp_root / "test.pdf")
    import pymupdf

    text = "".join(p.get_text() for p in pymupdf.open(str(pdf))).replace("\xa0", " ")
    assert "테스트 보고서" in text and "참고문헌" in text
    assert pdf.stat().st_size < 1_500_000  # fonts are subset
