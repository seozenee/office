import io

import openpyxl
import pymupdf
from docx import Document
from pptx import Presentation
from pptx.util import Inches

from app.knowledge.chunking import chunk_document, split_sentences
from app.knowledge.parsers import extract_references, find_date, parse_bytes


def _docx_bytes():
    d = Document()
    d.core_properties.author = "홍길동"
    d.add_heading("서론", level=1)
    d.add_paragraph("이 보고서는 2025년 4월 1일 작성되었다.")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "연도", "매출", "2024", "100"
    d.add_heading("참고문헌", level=1)
    d.add_paragraph("[1] 통계청 (2024). 인구 통계.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_parse_docx_sections_tables_references():
    pd = parse_bytes(_docx_bytes(), "a.docx")
    assert [s["heading"] for s in pd.sections][:1] == ["서론"]
    assert pd.tables and pd.tables[0]["rows"][1] == ["2024", "100"]
    assert pd.author == "홍길동"
    assert any("통계청" in r for r in pd.references)


def test_parse_pptx_slides_tables_notes():
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[5])
    s.shapes.title.text = "시장 현황"
    tbl = s.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(4), Inches(1)).table
    tbl.cell(0, 0).text = "A"
    tbl.cell(1, 1).text = "42"
    s.notes_slide.notes_text_frame.text = "발표자 노트 내용"
    buf = io.BytesIO()
    prs.save(buf)
    pd = parse_bytes(buf.getvalue(), "deck.pptx")
    assert pd.pages[0].section == "시장 현황"
    assert "발표자 노트 내용" in pd.text
    assert pd.tables[0]["rows"][-1][-1] == "42"


def test_parse_xlsx_sheets_as_tables():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Data"
    ws.append(["year", "value"])
    ws.append([2024, 10])
    buf = io.BytesIO()
    wb.save(buf)
    pd = parse_bytes(buf.getvalue(), "b.xlsx")
    assert pd.sections[0]["heading"] == "Data"
    assert pd.tables[0]["rows"][1] == ["2024", "10"]


def test_parse_pdf_pages_metadata_and_table():
    doc = pymupdf.open()
    for i in range(2):
        p = doc.new_page()
        p.insert_text((72, 72), f"Page {i + 1} heading\nMarket size reached 12 billion dollars in 2024.", fontsize=11)
    doc.set_metadata({"title": "Test PDF", "author": "Analyst", "creationDate": "D:20250102000000"})
    pd = parse_bytes(doc.tobytes(), "t.pdf")
    assert pd.title == "Test PDF" and pd.author == "Analyst" and pd.date == "2025-01-02"
    assert len(pd.pages) == 2 and pd.pages[1].number == 2
    assert "12 billion" in pd.pages[1].text


def test_parse_html_strips_nav_and_reads_dates():
    html = b"""<html><head><title>T</title><meta property="article:published_time" content="2025-02-03"></head>
    <body><nav>MENU ITEM</nav><article><h2>Intro</h2><p>Body text here.</p></article><script>evil()</script></body></html>"""
    pd = parse_bytes(html, "x.html")
    assert pd.date == "2025-02-03"
    assert "MENU ITEM" not in pd.text and "evil" not in pd.text
    assert pd.sections[0]["heading"] == "Intro"


def test_markdown_tables_and_headings():
    md = "# Title\n\n## Part\n\n| a | b |\n|---|---|\n| 1 | 2 |\n"
    pd = parse_bytes(md.encode(), "n.md")
    assert pd.title == "Title" and pd.tables[0]["rows"] == [["a", "b"], ["1", "2"]]


def test_dates_and_references_helpers():
    assert find_date("발표일: 2024년 3월 5일") == "2024-03-05"
    assert find_date("Published March 7, 2023") == "2023-03-07"
    refs = extract_references("본문\n\nReferences\n[1] Smith (2020). A study.\n[2] Kim (2021). Another.")
    assert len(refs) == 2


def test_sentence_split_rejoins_wrapped_pdf_lines():
    sents = split_sentences("시장 구조\n국내 시장에서 영상 판독 분야가 전체\n매출의 45%를 차지한다. 다음 문장이다.")
    assert "국내 시장에서 영상 판독 분야가 전체 매출의 45%를 차지한다." in sents
    assert sents[0] == "시장 구조"


def test_chunking_keeps_page_numbers():
    doc = pymupdf.open()
    for i in range(3):
        doc.new_page().insert_text((72, 72), f"Page {i + 1}. Some content sentence for page {i + 1}.")
    chunks = chunk_document(parse_bytes(doc.tobytes(), "p.pdf"))
    assert [c.page for c in chunks] == [1, 2, 3]
