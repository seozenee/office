"""Real .docx generation: cover, TOC, headings, labelled statements, clickable citations,
genuine Word footnotes, bibliography, header/footer with page numbers, tables, charts."""
from __future__ import annotations

import tempfile
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from xml.sax.saxutils import escape

from app.generators.charts import render_chart_png
from app.generators.content import LABEL_COLORS, LABEL_KO, Block, ReportContent, Statement

FONT = "Malgun Gothic"


class _Footnotes:
    def __init__(self) -> None:
        self.items: list[str] = []

    def add(self, text: str) -> int:
        self.items.append(text)
        return len(self.items)  # ids 1..n (-1/0 are separators)

    def xml(self) -> bytes:
        W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        parts = [f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><w:footnotes xmlns:w="{W}">',
                 '<w:footnote w:type="separator" w:id="-1"><w:p><w:r><w:separator/></w:r></w:p></w:footnote>',
                 '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:r><w:continuationSeparator/></w:r></w:p></w:footnote>']
        for i, text in enumerate(self.items, start=1):
            parts.append(f'<w:footnote w:id="{i}"><w:p><w:pPr><w:rPr><w:sz w:val="16"/></w:rPr></w:pPr>'
                         f'<w:r><w:rPr><w:vertAlign w:val="superscript"/></w:rPr><w:footnoteRef/></w:r>'
                         f'<w:r><w:rPr><w:sz w:val="16"/></w:rPr><w:t xml:space="preserve"> {escape(text[:900])}</w:t></w:r></w:p></w:footnote>')
        parts.append("</w:footnotes>")
        return "".join(parts).encode("utf-8")


def _set_font(style, size: float | None = None) -> None:
    style.font.name = FONT
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rfonts.set(qn(attr), FONT)
    if size:
        style.font.size = Pt(size)


def _field(paragraph, instr: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    for tag, attrs in (("w:fldChar", {"w:fldCharType": "begin"}), ("w:instrText", None),
                       ("w:fldChar", {"w:fldCharType": "separate"})):
        el = OxmlElement(tag)
        if attrs:
            for k, v in attrs.items():
                el.set(qn(k), v)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = instr
        run._r.append(el)
    if placeholder:
        t = OxmlElement("w:t")
        t.text = placeholder
        run._r.append(t)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


def _bookmark(paragraph, name: str, bid: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bid))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bid))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _internal_link(paragraph, text: str, anchor: str, superscript: bool = True) -> None:
    link = OxmlElement("w:hyperlink")
    link.set(qn("w:anchor"), anchor)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2B5DAA")
    rpr.append(color)
    if superscript:
        va = OxmlElement("w:vertAlign")
        va.set(qn("w:val"), "superscript")
        rpr.append(va)
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    link.append(r)
    paragraph._p.append(link)


def _external_link(paragraph, text: str, url: str) -> None:
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "2B5DAA")
    rpr.extend([color, u])
    r.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    t.set(qn("xml:space"), "preserve")
    r.append(t)
    link.append(r)
    paragraph._p.append(link)


def _footnote_ref(paragraph, fid: int) -> None:
    r = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    va = OxmlElement("w:vertAlign")
    va.set(qn("w:val"), "superscript")
    rpr.append(va)
    r.append(rpr)
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), str(fid))
    r.append(ref)
    paragraph._p.append(r)


def _shade(cell, hex_color: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


class DocxBuilder:
    def __init__(self, content: ReportContent, template_path: str | None = None) -> None:
        self.c = content
        self.doc = Document(template_path) if template_path and Path(template_path).exists() else Document()
        if template_path and Path(template_path).exists():  # keep styles/headers of the template, drop its body
            body = self.doc.element.body
            for el in list(body):
                if el.tag != qn("w:sectPr"):
                    body.remove(el)
        self.footnotes = _Footnotes()
        self._setup()

    def _setup(self) -> None:
        styles = self.doc.styles
        _set_font(styles["Normal"], 10.5)
        for name, size in (("Title", 26), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
            _set_font(styles[name], size)
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.left_margin = sec.right_margin = Cm(2.3)
        sec.top_margin = sec.bottom_margin = Cm(2.2)
        cp = self.doc.core_properties
        cp.title, cp.author, cp.subject = self.c.title, self.c.author, self.c.subtitle

    # -- statements -------------------------------------------------------
    def _statement(self, paragraph, s: Statement) -> None:
        if s.label:
            lr = paragraph.add_run(f"[{LABEL_KO.get(s.label, s.label)}] ")
            lr.bold = True
            lr.font.size = Pt(8.5)
            lr.font.color.rgb = RGBColor.from_string(LABEL_COLORS.get(s.label, "444444"))
        paragraph.add_run(s.text)
        for n in sorted(set(s.citations)):
            _internal_link(paragraph, f"[{n}]", f"ref_{n}")
        if s.footnote:
            _footnote_ref(paragraph, self.footnotes.add(s.footnote))

    def _block(self, b: Block) -> None:
        if b.type in ("paragraph", "callout"):
            for s in b.statements:
                p = self.doc.add_paragraph()
                if b.type == "callout":
                    p.paragraph_format.left_indent = Cm(0.6)
                self._statement(p, s)
        elif b.type == "bullets":
            for s in b.statements:
                self._statement(self.doc.add_paragraph(style="List Bullet"), s)
        elif b.type == "table" and b.rows:
            self._table(b.header, b.rows, b.caption)
        elif b.type == "chart" and b.chart:
            with tempfile.TemporaryDirectory() as td:
                png = render_chart_png(b.chart, Path(td) / "chart.png")
                self.doc.add_picture(str(png), width=Cm(15))
            self._caption(b.caption or b.chart.get("title", ""))
        elif b.type == "image" and b.image_path and Path(b.image_path).exists():
            self.doc.add_picture(b.image_path, width=Cm(15))
            self._caption(b.caption or "")

    def _caption(self, text: str) -> None:
        if text:
            p = self.doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.italic = True
            r.font.size = Pt(9)

    def _table(self, header: list[str], rows: list[list[str]], caption: str | None) -> None:
        ncols = max(len(header), *(len(r) for r in rows)) if rows else len(header)
        t = self.doc.add_table(rows=0, cols=ncols)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        if header:
            cells = t.add_row().cells
            for i, h in enumerate(header):
                cells[i].text = h
                cells[i].paragraphs[0].runs[0].bold = True
                _shade(cells[i], "E8EEF7")
        for r in rows:
            cells = t.add_row().cells
            for i in range(ncols):
                cells[i].text = str(r[i]) if i < len(r) else ""
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
        self._caption(caption or "")
        self.doc.add_paragraph()

    # -- document parts ---------------------------------------------------
    def _cover(self) -> None:
        for _ in range(6):
            self.doc.add_paragraph()
        p = self.doc.add_paragraph(style="Title")
        p.add_run(self.c.title)
        if self.c.subtitle:
            sp = self.doc.add_paragraph()
            r = sp.add_run(self.c.subtitle)
            r.font.size = Pt(14)
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        for _ in range(8):
            self.doc.add_paragraph()
        meta = self.doc.add_paragraph()
        meta.add_run(f"{self.c.organization or self.c.author}\n{self.c.date}").font.size = Pt(11)
        legend = self.doc.add_paragraph()
        legend.add_run("표기 규칙: ").bold = True
        for k, v in LABEL_KO.items():
            r = legend.add_run(f"[{v}] ")
            r.bold = True
            r.font.color.rgb = RGBColor.from_string(LABEL_COLORS[k])
        legend.add_run(" · [n] 은 참고문헌 번호(클릭 시 이동)").font.size = Pt(9)
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _toc(self) -> None:
        self.doc.add_heading("목차", level=1)
        p = self.doc.add_paragraph()
        lines = []
        if self.c.summary:
            lines.append("요약")
        lines += [("    " * (s.level - 1)) + s.heading for s in self.c.sections]
        lines.append("참고문헌")
        _field(p, 'TOC \\o "1-3" \\h \\z \\u', "\n".join(lines))
        note = self.doc.add_paragraph()
        r = note.add_run("(Word에서 목차를 우클릭 → 필드 업데이트 하면 페이지 번호가 표시됩니다.)")
        r.font.size = Pt(8)
        r.italic = True
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def _header_footer(self) -> None:
        sec = self.doc.sections[0]
        sec.different_first_page_header_footer = True
        hp = sec.header.paragraphs[0]
        hp.text = self.c.title
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for r in hp.runs:
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(0x88, 0x88, 0x88)
        fp = sec.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _field(fp, "PAGE", "1")
        fp.add_run(" / ")
        _field(fp, "NUMPAGES", "1")

    def _bibliography(self) -> None:
        self.doc.add_heading("참고문헌", level=1)
        if not self.c.bibliography:
            self.doc.add_paragraph("인용된 출처가 없습니다.")
            return
        for i, e in enumerate(self.c.bibliography):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.8)
            p.paragraph_format.first_line_indent = Cm(-0.8)
            p.add_run(f"[{e.number}] ").bold = True
            parts = [e.title]
            if e.publisher:
                parts.append(e.publisher)
            parts.append(e.date or "발행일 미상")
            p.add_run(". ".join(parts) + ". ")
            if e.url:
                _external_link(p, e.url, e.url)
            meta = [f"접근일 {e.accessed}" if e.accessed else None, f"p.{e.page}" if e.page else None,
                    f"유형: {e.tier_label}" if e.tier_label else None]
            meta_s = " · ".join(m for m in meta if m)
            if meta_s:
                p.add_run(f" ({meta_s})").font.size = Pt(8.5)
            _bookmark(p, f"ref_{e.number}", 1000 + i)

    def _qc(self) -> None:
        if not self.c.qc:
            return
        self.doc.add_heading("부록: 품질 검사(QC) 결과", level=1)
        rows = [[item["name"], "통과" if item["passed"] else "미통과", item.get("detail", "")] for item in self.c.qc.get("checks", [])]
        if rows:
            self._table(["검사 항목", "결과", "세부"], rows, "자동 품질 검사 체크리스트")
        stats = self.c.qc.get("stats", {})
        if stats:
            self._table(["지표", "값"], [[k, str(v)] for k, v in stats.items()], "검증 통계")

    def build(self, path: Path) -> Path:
        self._cover()
        self._toc()
        self._header_footer()
        if self.c.summary:
            self.doc.add_heading("요약", level=1)
            for s in self.c.summary:
                self._statement(self.doc.add_paragraph(style="List Bullet"), s)
        for sec in self.c.sections:
            self.doc.add_heading(sec.heading, level=min(sec.level, 3))
            for b in sec.blocks:
                self._block(b)
        self._bibliography()
        self._qc()
        if self.footnotes.items:
            part = Part(PackURI("/word/footnotes.xml"),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml",
                        self.footnotes.xml(), self.doc.part.package)
            self.doc.part.relate_to(part, RT.FOOTNOTES)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
        return path


def build_docx(content: ReportContent, path: Path, template_path: str | None = None) -> Path:
    return DocxBuilder(content, template_path).build(path)


__all__ = ["build_docx", "DocxBuilder", "WD_SECTION"]
