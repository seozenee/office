"""Real .xlsx generation (openpyxl) with live formulas, native charts, filters and a
formula evaluator used to cross-check spreadsheet formulas against the Python calc engine."""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.tools.calc import Assumption, FinancialModel, safe_eval

HEADER_FILL = PatternFill("solid", fgColor="1B2A49")
HEADER_FONT = Font(bold=True, color="FFFFFF")
KIND_FILL = {"ASSUMPTION": "FFF4E0", "USER INPUT": "E7F0FF", "SOURCE": "E6F4EA", "ESTIMATE": "F1E8FA", "FACT": "E6F4EA"}
THIN = Side(style="thin", color="D0D5DD")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MODEL_COLUMNS = [("year", "연도"), ("customers", "고객 수"), ("new_customers", "신규 고객"), ("arpu", "ARPU"),
                 ("revenue", "매출"), ("gross_profit", "매출총이익"), ("fixed_costs", "고정비"),
                 ("acquisition_cost", "고객획득비"), ("operating_income", "영업이익"), ("cumulative_income", "누적 영업이익")]


def _header(ws, row: int, values: list[str]) -> None:
    for j, v in enumerate(values, start=1):
        c = ws.cell(row=row, column=j, value=v)
        c.fill, c.font, c.border = HEADER_FILL, HEADER_FONT, BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autowidth(ws, max_width: int = 60) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None:
                widths[c.column] = max(widths.get(c.column, 0), min(len(str(c.value)) + 2, max_width))
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(w, 8)


class XlsxBuilder:
    def __init__(self, title: str) -> None:
        self.wb = openpyxl.Workbook()
        self.title = title
        self.summary = self.wb.active
        self.summary.title = "Summary"
        self.assumption_cells: dict[str, str] = {}

    def add_summary(self, lines: list[tuple[str, Any]], legend: bool = True) -> None:
        ws = self.summary
        ws["A1"] = self.title
        ws["A1"].font = Font(bold=True, size=16)
        r = 3
        for k, v in lines:
            ws.cell(row=r, column=1, value=k).font = Font(bold=True)
            ws.cell(row=r, column=2, value=v)
            r += 1
        if legend:
            r += 1
            ws.cell(row=r, column=1, value="표기 규칙").font = Font(bold=True)
            for kind, fill in KIND_FILL.items():
                r += 1
                c = ws.cell(row=r, column=1, value=kind)
                c.fill = PatternFill("solid", fgColor=fill)
        _autowidth(ws)

    def add_assumptions(self, assumptions: list[Assumption], citation_numbers: dict[int, int] | None = None) -> None:
        ws = self.wb.create_sheet("Assumptions")
        _header(ws, 1, ["Key", "항목", "값", "단위", "구분", "출처", "비고"])
        for i, a in enumerate(assumptions, start=2):
            ws.cell(row=i, column=1, value=a.key)
            ws.cell(row=i, column=2, value=a.label)
            vc = ws.cell(row=i, column=3, value=a.value)
            vc.number_format = "0.0%" if a.unit == "%" else "#,##0"
            ws.cell(row=i, column=4, value=a.unit)
            kc = ws.cell(row=i, column=5, value=a.kind)
            kc.fill = PatternFill("solid", fgColor=KIND_FILL.get(a.kind, "FFFFFF"))
            src = f"[{citation_numbers[a.source_id]}]" if citation_numbers and a.source_id in citation_numbers else ""
            ws.cell(row=i, column=6, value=src)
            ws.cell(row=i, column=7, value=a.note)
            for j in range(1, 8):
                ws.cell(row=i, column=j).border = BORDER
            self.assumption_cells[a.key] = f"Assumptions!$C${i}"
        ws.freeze_panes = "A2"
        _autowidth(ws)

    def add_financial_model(self, model: FinancialModel) -> None:
        ws = self.wb.create_sheet("Model")
        A = self.assumption_cells
        ws["A1"] = "재무 모델 (모든 값은 Assumptions 시트의 가정에서 수식으로 계산됨)"
        ws["A1"].font = Font(bold=True, size=12)
        _header(ws, 3, [label for _, label in MODEL_COLUMNS])
        first = 4
        for idx, row in enumerate(model.rows):
            r = first + idx
            p = r - 1
            ws.cell(row=r, column=1, value=row["year"])
            if idx == 0:
                f = {"B": f"={A['customers_y1']}", "C": f"=MAX(B{r},0)", "D": f"={A['arpu']}", "G": f"={A['fixed_costs_y1']}", "J": f"=I{r}"}
            else:
                f = {"B": f"=B{p}*(1+{A['customer_growth']})", "C": f"=MAX(B{r}-B{p}*(1-{A['churn']}),0)",
                     "D": f"=D{p}*(1+{A['price_growth']})", "G": f"=G{p}*(1+{A['fixed_cost_growth']})", "J": f"=J{p}+I{r}"}
            f.update({"E": f"=B{r}*D{r}", "F": f"=E{r}*{A['gross_margin']}", "H": f"=C{r}*{A['cac']}", "I": f"=F{r}-G{r}-H{r}"})
            for col, formula in f.items():
                c = ws[f"{col}{r}"]
                c.value = formula
                c.number_format = "#,##0"
            for j in range(1, 11):
                ws.cell(row=r, column=j).border = BORDER
        last = first + len(model.rows) - 1
        s = last + 2
        ws.cell(row=s, column=1, value="LTV").font = Font(bold=True)
        ws.cell(row=s, column=2, value=f"={A['arpu']}*{A['gross_margin']}*(1/{A['churn']})").number_format = "#,##0"
        ws.cell(row=s + 1, column=1, value="LTV/CAC").font = Font(bold=True)
        ws.cell(row=s + 1, column=2, value=f"=B{s}/{A['cac']}").number_format = "0.00"
        ws.cell(row=s + 2, column=1, value="CAC 회수기간(월)").font = Font(bold=True)
        ws.cell(row=s + 2, column=2, value=f"={A['cac']}/({A['arpu']}*{A['gross_margin']}/12)").number_format = "0.0"

        bar = BarChart()
        bar.title, bar.y_axis.title = "매출 / 영업이익", "원"
        bar.add_data(Reference(ws, min_col=5, min_row=3, max_row=last), titles_from_data=True)
        bar.add_data(Reference(ws, min_col=9, min_row=3, max_row=last), titles_from_data=True)
        bar.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
        bar.height, bar.width = 8, 16
        ws.add_chart(bar, f"L3")
        line = LineChart()
        line.title = "고객 수"
        line.add_data(Reference(ws, min_col=2, min_row=3, max_row=last), titles_from_data=True)
        line.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
        line.height, line.width = 8, 16
        ws.add_chart(line, "L20")
        ws.freeze_panes = "B4"
        _autowidth(ws, 18)

    def add_table_sheet(self, name: str, header: list[str], rows: list[list[Any]], widths: dict[str, int] | None = None) -> None:
        ws = self.wb.create_sheet(name[:31])
        _header(ws, 1, header)
        for i, r in enumerate(rows, start=2):
            for j, v in enumerate(r, start=1):
                c = ws.cell(row=i, column=j, value=v)
                c.border = BORDER
                c.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
        if rows:
            ws.auto_filter.ref = f"A1:{get_column_letter(len(header))}{len(rows) + 1}"
        ws.freeze_panes = "A2"
        _autowidth(ws)
        for col, w in (widths or {}).items():
            ws.column_dimensions[col].width = w

    def save(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.wb.save(str(path))
        return path


# ---------------------------------------------------------------------------
# Formula evaluation (verification: spreadsheet formulas must equal Python results)
# ---------------------------------------------------------------------------
_REF = re.compile(r"(?:(?P<sheet>[A-Za-z_][A-Za-z0-9_]*)!)?\$?(?P<col>[A-Z]{1,3})\$?(?P<row>\d+)")


def evaluate_workbook(path: Path) -> dict[str, dict[str, float]]:
    """Evaluate arithmetic/MAX/MIN formulas in a workbook we generated. Returns {sheet: {cell: value}}."""
    wb = openpyxl.load_workbook(str(path), data_only=False)
    cache: dict[tuple[str, str], float] = {}

    def value(sheet: str, cell: str, depth: int = 0) -> float:
        key = (sheet, cell)
        if key in cache:
            return cache[key]
        if depth > 200:
            raise RecursionError("formula nesting too deep")
        raw = wb[sheet][cell].value
        if isinstance(raw, str) and raw.startswith("="):
            expr_vars: dict[str, float] = {}

            def repl(m: re.Match[str]) -> str:
                sh = m.group("sheet") or sheet
                name = f"v_{sh}_{m.group('col')}{m.group('row')}"
                expr_vars[name] = value(sh, f"{m.group('col')}{m.group('row')}", depth + 1)
                return name

            expr = _REF.sub(repl, raw[1:]).replace("MAX(", "max(").replace("MIN(", "min(")
            result = safe_eval(expr, expr_vars)
        elif raw is None:
            result = 0.0
        else:
            result = float(raw)
        cache[key] = result
        return result

    out: dict[str, dict[str, float]] = {}
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str) and c.value.startswith("="):
                    out.setdefault(ws.title, {})[c.coordinate] = value(ws.title, c.coordinate)
    return out


def verify_model_sheet(path: Path, model: FinancialModel, tolerance: float = 0.5) -> list[str]:
    """Compare every Model-sheet formula with the Python calc engine. Returns mismatch messages."""
    values = evaluate_workbook(path).get("Model", {})
    problems: list[str] = []
    for idx, row in enumerate(model.rows):
        r = 4 + idx
        for j, (key, _) in enumerate(MODEL_COLUMNS[1:], start=2):
            cell = f"{get_column_letter(j)}{r}"
            if cell in values and abs(values[cell] - row[key]) > tolerance + abs(row[key]) * 1e-6:
                problems.append(f"{cell} ({key}, year {row['year']}): formula={values[cell]:.2f} python={row[key]:.2f}")
    return problems
