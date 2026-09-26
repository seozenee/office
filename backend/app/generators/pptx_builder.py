"""Real .pptx generation with native charts, tables, timelines, process diagrams, comparison
tables, KPI cards, citations and speaker notes. Default storyline (spec §11):
Problem → Evidence → Insight → Solution → Business Model → Market → Strategy → Financials → Roadmap → Conclusion.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Literal

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

from app.generators.content import LABEL_COLORS, LABEL_KO, BibEntry, Statement, cite_marks

STORYLINE = ["Problem", "Evidence", "Insight", "Solution", "Business Model", "Market", "Strategy",
             "Financials", "Roadmap", "Conclusion"]

NAVY = RGBColor(0x1B, 0x2A, 0x49)
ACCENT = RGBColor(0x2B, 0x5D, 0xAA)
ORANGE = RGBColor(0xE0, 0x7B, 0x39)
LIGHT = RGBColor(0xF3, 0xF5, 0xF9)
GREY = RGBColor(0x66, 0x66, 0x66)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
FONT = "Malgun Gothic"


@dataclass
class SlideSpec:
    layout: Literal["title", "section", "bullets", "chart", "table", "timeline", "process", "comparison",
                    "kpi", "sources", "two_column"]
    title: str
    subtitle: str = ""
    kicker: str = ""  # small storyline label, e.g. "PROBLEM"
    statements: list[Statement] = field(default_factory=list)
    right_statements: list[Statement] = field(default_factory=list)
    chart: dict[str, Any] | None = None
    header: list[str] = field(default_factory=list)
    rows: list[list[str]] = field(default_factory=list)
    steps: list[dict[str, str]] = field(default_factory=list)  # timeline/process: {"label","text"}
    kpis: list[dict[str, str]] = field(default_factory=list)  # {"value","label","note"}
    bibliography: list[BibEntry] = field(default_factory=list)
    notes: str = ""


def _txt(tf, text: str, size: float, *, bold: bool = False, color: RGBColor = NAVY, align=PP_ALIGN.LEFT) -> None:
    tf.clear()
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
    r.font.color.rgb = color


class PptxBuilder:
    W, H = Inches(13.333), Inches(7.5)

    def __init__(self, title: str, footer: str = "") -> None:
        self.prs = Presentation()
        self.prs.slide_width, self.prs.slide_height = self.W, self.H
        self.prs.core_properties.title = title
        self.footer = footer
        self.blank = self.prs.slide_layouts[6]
        self.count = 0

    # -- primitives -------------------------------------------------------
    def _box(self, slide, x, y, w, h, fill: RGBColor | None = None, shape=MSO_SHAPE.RECTANGLE, line: RGBColor | None = None):
        s = slide.shapes.add_shape(shape, x, y, w, h)
        if fill is None:
            s.fill.background()
        else:
            s.fill.solid()
            s.fill.fore_color.rgb = fill
        if line is None:
            s.line.fill.background()
        else:
            s.line.color.rgb = line
        s.shadow.inherit = False
        return s

    def _text(self, slide, x, y, w, h, text: str, size: float, **kw):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tb.text_frame.word_wrap = True
        _txt(tb.text_frame, text, size, **kw)
        return tb

    def _frame(self, spec: SlideSpec):
        self.count += 1
        slide = self.prs.slides.add_slide(self.blank)
        self._box(slide, 0, 0, self.W, Inches(0.12), ACCENT)
        if spec.kicker:
            self._text(slide, Inches(0.6), Inches(0.35), Inches(8), Inches(0.35), spec.kicker.upper(), 11, bold=True, color=ORANGE)
        self._text(slide, Inches(0.6), Inches(0.65), Inches(12), Inches(0.9), spec.title, 28, bold=True)
        foot = f"{self.footer}   |   {self.count}"
        self._text(slide, Inches(0.6), Inches(7.0), Inches(12.1), Inches(0.3), foot, 9, color=GREY, align=PP_ALIGN.RIGHT)
        if spec.notes:
            slide.notes_slide.notes_text_frame.text = spec.notes
        return slide

    def _statements(self, slide, stmts: list[Statement], x, y, w, h, size: float = 16) -> None:
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        first = True
        for s in stmts:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.space_after = Pt(10)
            if s.label:
                lr = p.add_run()
                lr.text = f"{LABEL_KO.get(s.label, s.label)}  "
                lr.font.size, lr.font.bold, lr.font.name = Pt(size * 0.65), True, FONT
                lr.font.color.rgb = RGBColor.from_string(LABEL_COLORS.get(s.label, "444444"))
            r = p.add_run()
            r.text = "▪ " + s.text if not s.label else s.text
            r.font.size, r.font.name = Pt(size), FONT
            r.font.color.rgb = NAVY
            if s.citations:
                c = p.add_run()
                c.text = " " + cite_marks(s.citations)
                c.font.size, c.font.name = Pt(size * 0.6), FONT
                c.font.color.rgb = ACCENT
                c.font._rPr.set("baseline", "30000")

    # -- layouts ----------------------------------------------------------
    def title_slide(self, spec: SlideSpec) -> None:
        self.count += 1
        slide = self.prs.slides.add_slide(self.blank)
        self._box(slide, 0, 0, self.W, self.H, NAVY)
        self._box(slide, Inches(0.8), Inches(3.55), Inches(1.4), Inches(0.08), ORANGE)
        self._text(slide, Inches(0.8), Inches(1.9), Inches(11.5), Inches(1.6), spec.title, 40, bold=True, color=WHITE)
        self._text(slide, Inches(0.8), Inches(3.8), Inches(11.5), Inches(1.0), spec.subtitle, 18, color=RGBColor(0xC8, 0xD3, 0xE6))
        if spec.notes:
            slide.notes_slide.notes_text_frame.text = spec.notes

    def section_slide(self, spec: SlideSpec) -> None:
        self.count += 1
        slide = self.prs.slides.add_slide(self.blank)
        self._box(slide, 0, 0, Inches(4.2), self.H, ACCENT)
        self._text(slide, Inches(0.6), Inches(3.0), Inches(3.4), Inches(1.2), spec.kicker or "", 16, bold=True, color=WHITE)
        self._text(slide, Inches(4.8), Inches(2.8), Inches(8), Inches(1.6), spec.title, 34, bold=True)
        if spec.notes:
            slide.notes_slide.notes_text_frame.text = spec.notes

    def bullets(self, spec: SlideSpec) -> None:
        slide = self._frame(spec)
        self._statements(slide, spec.statements, Inches(0.8), Inches(1.7), Inches(11.8), Inches(5.1))

    def two_column(self, spec: SlideSpec) -> None:
        slide = self._frame(spec)
        self._box(slide, Inches(0.6), Inches(1.7), Inches(5.9), Inches(5.0), LIGHT)
        self._box(slide, Inches(6.8), Inches(1.7), Inches(5.9), Inches(5.0), LIGHT)
        self._statements(slide, spec.statements, Inches(0.8), Inches(1.9), Inches(5.5), Inches(4.6), 14)
        self._statements(slide, spec.right_statements, Inches(7.0), Inches(1.9), Inches(5.5), Inches(4.6), 14)

    def chart(self, spec: SlideSpec) -> None:
        slide = self._frame(spec)
        c = spec.chart or {}
        data = CategoryChartData()
        data.categories = [str(x) for x in c.get("categories", [])]
        for s in c.get("series", []):
            data.add_series(s["name"], [float(v) for v in s["values"]])
        kind = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
                "pie": XL_CHART_TYPE.PIE}.get(c.get("kind", "bar"), XL_CHART_TYPE.COLUMN_CLUSTERED)
        width = Inches(7.8) if spec.statements else Inches(12)
        gf = slide.shapes.add_chart(kind, Inches(0.6), Inches(1.7), width, Inches(5.0), data)
        ch = gf.chart
        ch.has_legend = len(c.get("series", [])) > 1 or kind == XL_CHART_TYPE.PIE
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
        if kind != XL_CHART_TYPE.PIE:
            ch.value_axis.has_major_gridlines = True
            ch.value_axis.tick_labels.number_format = c.get("number_format", "#,##0")
            ch.value_axis.tick_labels.number_format_is_linked = False
            for i, plot_series in enumerate(ch.plots[0].series):
                fmt = plot_series.format
                color = [ACCENT, ORANGE, RGBColor(0x1F, 0x7A, 0x4D)][i % 3]
                if kind == XL_CHART_TYPE.LINE_MARKERS:
                    fmt.line.color.rgb = color
                else:
                    fmt.fill.solid()
                    fmt.fill.fore_color.rgb = color
        if spec.statements:
            self._statements(slide, spec.statements, Inches(8.7), Inches(1.9), Inches(4.2), Inches(4.8), 13)
        if c.get("source_note"):
            self._text(slide, Inches(0.6), Inches(6.7), Inches(12), Inches(0.3), c["source_note"], 9, color=GREY)

    def table(self, spec: SlideSpec) -> None:
        slide = self._frame(spec)
        rows = spec.rows[:12]
        ncols = max(len(spec.header), *(len(r) for r in rows)) if rows else len(spec.header)
        nrows = len(rows) + (1 if spec.header else 0)
        if not nrows or not ncols:
            return
        shape = slide.shapes.add_table(nrows, ncols, Inches(0.6), Inches(1.7), Inches(12.1), Inches(0.4) * nrows)
        tbl = shape.table
        r0 = 0
        if spec.header:
            for j, h in enumerate(spec.header):
                cell = tbl.cell(0, j)
                _txt(cell.text_frame, h, 12, bold=True, color=WHITE)
                cell.fill.solid()
                cell.fill.fore_color.rgb = NAVY
            r0 = 1
        for i, r in enumerate(rows):
            for j in range(ncols):
                cell = tbl.cell(i + r0, j)
                _txt(cell.text_frame, str(r[j]) if j < len(r) else "", 11)
                cell.fill.solid()
                cell.fill.fore_color.rgb = LIGHT if i % 2 == 0 else WHITE
        if spec.statements:
            self._statements(slide, spec.statements, Inches(0.6), Inches(1.9) + Inches(0.4) * nrows, Inches(12), Inches(1.5), 12)

    def comparison(self, spec: SlideSpec) -> None:
        self.table(spec)

    def timeline(self, spec: SlideSpec) -> None:
        slide = self._frame(spec)
        steps = spec.steps[:6] or [{"label": "-", "text": "UNKNOWN"}]
        n = len(steps)
        y = Inches(3.6)
        self._box(slide, Inches(0.8), y, Inches(11.7), Inches(0.06), ACCENT)
        span = Inches(11.7) / n
        for i, st in enumerate(steps):
            cx = Inches(0.8) + span * i + span / 2
            self._box(slide, Emu(int(cx - Inches(0.18))), Emu(int(y - Inches(0.15))), Inches(0.36), Inches(0.36), ORANGE, MSO_SHAPE.OVAL)
            self._text(slide, Emu(int(cx - span / 2)), Inches(2.6), Emu(int(span)), Inches(0.8), st["label"], 15, bold=True, align=PP_ALIGN.CENTER)
            self._text(slide, Emu(int(cx - span / 2 + Inches(0.1))), Inches(4.2), Emu(int(span - Inches(0.2))), Inches(2.4), st["text"], 12, color=GREY, align=PP_ALIGN.CENTER)

    def process(self, spec: SlideSpec) -> None:
        slide = self._frame(spec)
        steps = spec.steps[:6] or [{"label": "-", "text": "UNKNOWN"}]
        n = len(steps)
        w = Inches(12) / n
        for i, st in enumerate(steps):
            x = Inches(0.6) + w * i
            shp = self._box(slide, Emu(int(x)), Inches(2.2), Emu(int(w - Inches(0.05))), Inches(1.2),
                            ACCENT if i % 2 == 0 else NAVY, MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON)
            _txt(shp.text_frame, st["label"], 14, bold=True, color=WHITE, align=PP_ALIGN.CENTER)
            shp.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
            self._text(slide, Emu(int(x + Inches(0.1))), Inches(3.7), Emu(int(w - Inches(0.3))), Inches(2.8), st["text"], 12, color=GREY)

    def kpi(self, spec: SlideSpec) -> None:
        slide = self._frame(spec)
        kpis = spec.kpis[:4]
        n = max(len(kpis), 1)
        w = Inches(12) / n
        for i, k in enumerate(kpis):
            x = Emu(int(Inches(0.6) + w * i))
            self._box(slide, x, Inches(1.9), Emu(int(w - Inches(0.25))), Inches(2.6), LIGHT)
            self._text(slide, x, Inches(2.2), Emu(int(w - Inches(0.25))), Inches(1.0), k["value"], 30, bold=True, color=ACCENT, align=PP_ALIGN.CENTER)
            self._text(slide, x, Inches(3.3), Emu(int(w - Inches(0.25))), Inches(0.6), k["label"], 14, bold=True, align=PP_ALIGN.CENTER)
            if k.get("note"):
                self._text(slide, x, Inches(3.85), Emu(int(w - Inches(0.25))), Inches(0.6), k["note"], 10, color=GREY, align=PP_ALIGN.CENTER)
        if spec.statements:
            self._statements(slide, spec.statements, Inches(0.8), Inches(4.8), Inches(11.8), Inches(2.0), 13)

    def sources(self, spec: SlideSpec) -> None:
        entries = spec.bibliography or []
        per = 9
        for start in range(0, max(len(entries), 1), per):
            page = entries[start:start + per]
            s = SlideSpec("bullets", spec.title if start == 0 else f"{spec.title} (계속)", kicker="SOURCES", notes=spec.notes)
            slide = self._frame(s)
            tb = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.3))
            tf = tb.text_frame
            tf.word_wrap = True
            if not page:
                _txt(tf, "인용된 출처가 없습니다.", 12)
            for i, e in enumerate(page):
                p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                r = p.add_run()
                r.text = f"[{e.number}] {e.title} — {e.publisher or ''} ({e.date or '발행일 미상'}) "
                r.font.size, r.font.name = Pt(11), FONT
                if e.url:
                    lr = p.add_run()
                    lr.text = e.url
                    lr.font.size, lr.font.name = Pt(9), FONT
                    lr.font.color.rgb = ACCENT
                    lr.hyperlink.address = e.url

    def add(self, spec: SlideSpec) -> None:
        getattr(self, {"title": "title_slide", "section": "section_slide", "bullets": "bullets", "chart": "chart",
                       "table": "table", "timeline": "timeline", "process": "process", "comparison": "comparison",
                       "kpi": "kpi", "sources": "sources", "two_column": "two_column"}[spec.layout])(spec)

    def save(self, path: Path) -> Path:
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(str(path))
        return path


def build_pptx(title: str, slides: list[SlideSpec], path: Path, footer: str = "") -> Path:
    b = PptxBuilder(title, footer or title)
    for s in slides:
        b.add(s)
    return b.save(path)
